import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_PRIMARY = RGBColor(43, 92, 146)    # #2B5C92 (Navy Pekat)
    COLOR_DARK = RGBColor(12, 20, 70)       # #0C1446 (Midnight Dark Navy)
    COLOR_ACCENT = RGBColor(255, 216, 0)    # #FFD800 (Bumblebee Yellow)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_GRAY_BG = RGBColor(241, 245, 249) # #F1F5F9 (Light Slate)
    COLOR_TEXT_MAIN = RGBColor(30, 41, 59)  # #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B
    COLOR_CARD_BG = RGBColor(255, 255, 255)
    COLOR_BORDER = RGBColor(226, 232, 240)

    blank_slide_layout = prs.slide_layouts[6] # Blank layout

    def add_header(slide, title_text, category_text="BEASISWA BLITAR MENGABDI V2.2.0"):
        # Top banner background
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_PRIMARY
        top_bar.line.fill.background()

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(12), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        # Title text
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE

        # Bottom footer line
        footer_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.1), Inches(13.333), Inches(0.4))
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = COLOR_DARK
        footer_line.line.fill.background()

        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.15), Inches(12.133), Inches(0.3))
        tf_foot = footer_box.text_frame
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "Dinas Kepemudaan dan Olahraga (Dispora) Kabupaten Blitar  |  Dokumen Presentasi Resmi V2.2.0"
        p_foot.font.name = "Arial"
        p_foot.font.size = Pt(9)
        p_foot.font.color.rgb = COLOR_WHITE

    def add_card(slide, left, top, width, height, title, items, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        # Title box
        t_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.5))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY

        # Body items
        b_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.65), width - Inches(0.4), height - Inches(0.8))
        tf_b = b_box.text_frame
        tf_b.word_wrap = True
        
        for i, item in enumerate(items):
            p = tf_b.paragraphs[0] if i == 0 else tf_b.add_paragraph()
            p.text = f"•  {item}"
            p.font.name = "Arial"
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_after = Pt(6)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1: COVER SLIDE
    # ═════════════════════════════════════════════════════════════════════════
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_DARK
    bg1.line.fill.background()

    # Decorative Accent Bar
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(0.2), Inches(4.5))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_ACCENT
    accent_bar.line.fill.background()

    # Title Box
    t_box1 = slide1.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.5), Inches(3.5))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "BEASISWA BLITAR MENGABDI"
    p1.font.name = "Arial"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ACCENT

    p2 = tf1.add_paragraph()
    p2.text = "Sistem Informasi Manajemen Beasiswa Terpadu V2.2.0"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "Presentasi Overview Fitur, Arsitektur Keamanan, Sistem Pendukung Keputusan (SPK), dan Workflow Berjenjang"
    p3.font.name = "Arial"
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_MUTED
    p3.space_before = Pt(15)

    p4 = tf1.add_paragraph()
    p4.text = "Pemerintah Kabupaten Blitar — Dinas Kepemudaan dan Olahraga (Dispora)  |  2026"
    p4.font.name = "Arial"
    p4.font.size = Pt(11)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_WHITE
    p4.space_before = Pt(30)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2: LATAR BELAKANG & TUJUAN
    # ═════════════════════════════════════════════════════════════════════════
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Latar Belakang & Pilar Utama Digitalisasi Beasiswa")

    add_card(slide2, Inches(0.6), Inches(1.5), Inches(3.7), Inches(5.2), "Peralihan Digital", [
        "Digitalisasi penuh dari pendaftaran fisik/kertas ke portal web terpadu.",
        "Mengurangi risiko penumpukan berkas dan mempercepat proses verifikasi.",
        "Memudahkan pemantauan progres pendaftaran secara real-time."
    ])

    add_card(slide2, Inches(4.8), Inches(1.5), Inches(3.7), Inches(5.2), "Objektivitas SPK", [
        "Perhitungan skor otomatis berbasis Weighted Sum Model (WSM).",
        "Menghilangkan preferensi subjektif dalam menentukan urutan peringkat.",
        "Aturan tie-breaker transparan berdasarkan waktu pendaftaran awal."
    ])

    add_card(slide2, Inches(9.0), Inches(1.5), Inches(3.7), Inches(5.2), "Akses Pendaftaran Publik", [
        "Pendaftaran mandiri tanpa perlu membuat akun pengguna (Guest 5-step).",
        "Pemeriksaan status kelulusan cukup menggunakan NIK 16 digit & Tahun.",
        "Unduh langsung Tanda Bukti Pendaftaran PDF berstempel resmi."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3: TIGA PROGRAM BEASISWA UTAMA
    # ═════════════════════════════════════════════════════════════════════════
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Tiga Program Beasiswa Utama Kabupaten Blitar")

    add_card(slide3, Inches(0.6), Inches(1.5), Inches(3.7), Inches(5.2), "Satu Desa Satu Sarjana (SDSS)", [
        "Program beasiswa penuh perwakilan desa/kelurahan di Kab. Blitar.",
        "Alur verifikasi khusus: Faktual Desa (pilih 1 calon) & peninjauan Kecamatan.",
        "Aturan ketat: 1 NIK hanya berlaku 1 kali seumur hidup."
    ])

    add_card(slide3, Inches(4.8), Inches(1.5), Inches(3.7), Inches(5.2), "Berdaya Berjaya", [
        "Program beasiswa prestasi untuk mahasiswa Baru & Lama.",
        "Dilengkapi dengan tahapan Tes Wawancara oleh tim penilai.",
        "Nilai wawancara terintegrasi langsung ke dalam bobot kriteria SPK."
    ])

    add_card(slide3, Inches(9.0), Inches(1.5), Inches(3.7), Inches(5.2), "Bantuan Biaya Pendidikan (BBP)", [
        "Bantuan biaya pendidikan mahasiswa kurang mampu / berprestasi.",
        "Jalur Khusus: Jalur Prestasi & Jalur Kurang Mampu (KIP/PKH).",
        "Seleksi cepat berbasis pembobotan kriteria ekonomi & akademik."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4: PANDUAN PENDAFTARAN PUBLIK (5-STEP)
    # ═════════════════════════════════════════════════════════════════════════
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "Formulir Pendaftaran Online Mandiri (5-Step Wizard)")

    add_card(slide4, Inches(0.6), Inches(1.5), Inches(5.8), Inches(5.2), "Alur Pengisian Formulir 5-Step", [
        "Step 1 — Identitas: NIK 16 digit, Nama, Alamat, No. WA, Kampus & Prodi.",
        "Step 2 — Orang Tua/Wali: Data Ayah & Ibu, NIK Orang Tua, Pekerjaan, Penghasilan.",
        "Step 3 — Ekonomi: Penghasilan Gabungan, Rumah/Lahan, KIP/PKH.",
        "Step 4 — Akademik: Nilai IPK / Rapor, Prestasi Juara, Keaktifan Organisasi.",
        "Step 5 — Dokumen: Upload KTP, KK, Surat Permohonan, Pernyataan, SPTJM."
    ])

    add_card(slide4, Inches(6.8), Inches(1.5), Inches(5.8), Inches(5.2), "Fitur Review & Cek Status NIK", [
        "Review Pendaftaran: Memeriksa ringkasan seluruh data sebelum disubmit.",
        "Persetujuan Privasi: Menyetujui keabsahan data & kebijakan privasi.",
        "Penguncian Permanen: Data & berkas yang disubmit terkunci permanen.",
        "Cek Status NIK: Memantau kelulusan via 16 digit NIK & Tahun.",
        "Cetak Bukti PDF: Mengunduh Tanda Bukti Pendaftaran PDF berstempel."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5: ARSITEKTUR KEAMANAN & PROTEKSI
    # ═════════════════════════════════════════════════════════════════════════
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Arsitektur Keamanan & Proteksi Sistem V2.2.0")

    add_card(slide5, Inches(0.6), Inches(1.5), Inches(3.7), Inches(5.2), "URL Login Terselubung", [
        "Route login diubah menjadi '/mengabdi' (tanpa tombol di publik).",
        "Admin wajib mengetik URL secara manual pada address bar.",
        "Mencegah pencarian otomatis (enumeration) halaman admin oleh bot."
    ])

    add_card(slide5, Inches(4.8), Inches(1.5), Inches(3.7), Inches(5.2), "Penguncian NIK Permanen", [
        "Setiap pendaftaran terkunci berdasarkan NIK 16 digit resmi.",
        "Mencegah pendaftaran ganda pada periode pendaftaran aktif.",
        "Merekan histori penerima tahun lalu untuk mencegah penerimaan ganda."
    ])

    add_card(slide5, Inches(9.0), Inches(1.5), Inches(3.7), Inches(5.2), "Throttle & Audit Log", [
        "Pembatasan percobaan login (max 5x/menit) untuk cegah brute force.",
        "Audit Log mencatat IP Address, User, Timestamp, & rincian aksi.",
        "Penutupan akses mutasi data publik setelah tombol Submit diklik."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6: STRUKTUR MULTI-ROLE USER (5 HAK AKSES)
    # ═════════════════════════════════════════════════════════════════════════
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "Struktur Hak Akses Administrator Berjenjang (5 Roles)")

    add_card(slide6, Inches(0.6), Inches(1.5), Inches(5.8), Inches(2.5), "1. Super Admin (/super-admin)", [
        "Akses penuh master data, pengguna, referensi wilayah, & perbankan.",
        "Konfigurasi Periode, Program, Jalur, Persyaratan, Kriteria & Bobot SPK.",
        "Mengelola Dokumen Publik (Ebook/Juknis) & memantau Audit Log System."
    ])

    add_card(slide6, Inches(6.8), Inches(1.5), Inches(5.8), Inches(2.5), "2. Admin Kabupaten (/kabupaten)", [
        "Verifikasi administrasi akhir & menginput skor Tes Wawancara.",
        "Menjalankan pemeringkatan SPK & menetapkan SK Bupati Terbit.",
        "Menerbitkan Output Administrasi & pencatatan data penyaluran pembayaran."
    ])

    add_card(slide6, Inches(0.6), Inches(4.2), Inches(3.7), Inches(2.5), "3. Admin OPD (/opd)", [
        "Verifikasi keabsahan dokumen sesuai bidang instansi kewenangan.",
        "Status: Valid, Tidak Valid, atau Tidak Wajib + Catatan Alasan."
    ])

    add_card(slide6, Inches(4.8), Inches(4.2), Inches(3.7), Inches(2.5), "4. Admin Kecamatan (/kecamatan)", [
        "Meninjau rekomendasi pendaftar SDSS tingkat desa.",
        "Validasi peninjauan Kecamatan & meneruskan ke Kabupaten."
    ])

    add_card(slide6, Inches(9.0), Inches(4.2), Inches(3.7), Inches(2.5), "5. Admin Desa (/desa)", [
        "Verifikasi faktual domisili pendaftar SDSS.",
        "Melihat skor SPK desa, memilih 1 calon & upload Surat Rekomendasi."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7: SPK WEIGHTED SUM & TIE-BREAKER
    # ═════════════════════════════════════════════════════════════════════════
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "Sistem Pendukung Keputusan (SPK) & Aturan Tie-Breaker")

    add_card(slide7, Inches(0.6), Inches(1.5), Inches(5.8), Inches(5.2), "Weighted Sum Model (WSM)", [
        "Penilaian otomatis berbasis pembobotan kriteria terkonfigurasi.",
        "Formula: Total Skor = Sum (Skor Pilihan / Skor Maks * Bobot Kriteria).",
        "Kriteria mencakup: Ekonomi, IPK, Prestasi, Organisasi, Wawancara.",
        "Skor terhitung secara real-time dan dapat dire-kalkulasi sewaktu-waktu.",
        "Menghasilkan urutan pemeringkatan objektif dari nilai tertinggi."
    ])

    add_card(slide7, Inches(6.8), Inches(1.5), Inches(5.8), Inches(5.2), "Aturan Penentu Nilai Sama (Tie-Breaker)", [
        "Kondisi: Apabila terdapat 2 atau lebih pemohon dengan Total Skor SPK persis.",
        "Solusi Otomatis: Sistem memeriksa timestamp waktu submit pendaftaran.",
        "Prioritas: Pemohon yang mendaftar LEBIH AWAL diposisikan lebih tinggi.",
        "Transparansi: Tidak memerlukan keputusan manual subjektif dari panitia.",
        "Auditable: Seluruh log waktu tercatat presisi hingga detik."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 8: END-TO-END WORKFLOW
    # ═════════════════════════════════════════════════════════════════════════
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "Alur Kerja Terpadu (End-to-End Workflow)")

    flow_box = slide8.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.133), Inches(5.2))
    tf_f = flow_box.text_frame
    tf_f.word_wrap = True

    steps = [
        ("1. Pengajuan Publik", "Pemohon mengisi formulir 5-step & melakukan review submit mandiri."),
        ("2. Verifikasi OPD", "Admin OPD memeriksa fisik dokumen persyaratan & memberi status keabsahan."),
        ("3. Faktual Desa & Kec (SDSS)", "Desa memilih 1 calon terbaik, mengunggah rekomendasi, & Kecamatan menyetujui."),
        ("4. Tes Wawancara (Berdaya)", "Admin Kabupaten menginput skor tes wawancara pemohon ke dalam sistem."),
        ("5. Pemeringkatan SPK", "Mesin SPK menghitung skor terbobot Weighted Sum & menerapkan tie-breaker."),
        ("6. Penetapan SK Terbit", "Admin Kabupaten menetapkan kelulusan 'Lulus — SK Terbit' & data terkunci final."),
        ("7. Output & Pembayaran", "Penerbitan Draft SK Bupati, Rekapitulasi Output, & pencatatan penyaluran dana.")
    ]

    for title, desc in steps:
        p = tf_f.add_paragraph() if tf_f.paragraphs[0].text else tf_f.paragraphs[0]
        p.text = f"{title}: {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.space_after = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 9: DOKUMEN PUBLIK & OUTPUT ADMINISTRASI
    # ═════════════════════════════════════════════════════════════════════════
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "Dokumen Publik & Modul Output Administrasi")

    add_card(slide9, Inches(0.6), Inches(1.5), Inches(5.8), Inches(5.2), "Dokumen Publik (Front-End)", [
        "Ebook Panduan Pendaftaran 2026 (.pdf) untuk petunjuk teknis pemohon.",
        "SK Juknis Beasiswa Blitar Mengabdi (.pdf) dari Kepala Dinas Dispora.",
        "Template Surat Permohonan resmi (.docx) untuk 3 program beasiswa.",
        "Template Surat Pernyataan & SPTJM Kebenaran Data (.docx) bermaterai.",
        "Seluruh dokumen dapat diunduh bebas oleh publik di halaman Informasi."
    ])

    add_card(slide9, Inches(6.8), Inches(1.5), Inches(5.8), Inches(5.2), "Output Administrasi (Back-End)", [
        "Draft SK Bupati: Generasi naskah Keputusan Bupati Blitar daftar penerima.",
        "Rekapitulasi Penerima: Laporan rekap per kecamatan, desa, & perguruan tinggi.",
        "Pencatatan Pembayaran: Data nominal, tanggal pencairan, & status transfer.",
        "Ekspor Data Excel/PDF: Laporan siap cetak untuk kebutuhan auditisasi."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 10: HASIL PENGUJIAN & KUALITAS KODE
    # ═════════════════════════════════════════════════════════════════════════
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide10, "Kualitas Kode & Hasil Pengujian Sistem (Automated Tests)")

    add_card(slide10, Inches(0.6), Inches(1.5), Inches(5.8), Inches(5.2), "Hasil Automated Testing", [
        "Total Test Suite: 27 Tests Passed (81 Assertions - 100% Success).",
        "Pengujian Auth: Login, Logout, Ganti Password, & Proteksi Unauthorized.",
        "Pengujian NIK Locking: Mencegah NIK terdaftar 2x & memvalidasi histori.",
        "Pengujian Dokumen Publik: Upload, Manajemen Super Admin, & Unduh Publik.",
        "Pengujian End-to-End: Pengujian siklus penuh beasiswa dari submit ke SK."
    ])

    add_card(slide10, Inches(6.8), Inches(1.5), Inches(5.8), Inches(5.2), "Arsitektur Clean Code", [
        "Framework: Laravel 12 dengan Thin Controller + Service Layer pattern.",
        "Frontend: Tailwind CSS v4 + Alpine.js + Lucide Icons (Full Responsive).",
        "Seeder Teruji: DokumenPublikSeeder membaca file fisik asli .docx.",
        "Database Rules: Foreign Key constraints & Indexing ter-optimasi."
    ])

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 11: KESIMPULAN & CONTACT SUPPORT
    # ═════════════════════════════════════════════════════════════════════════
    slide11 = prs.slides.add_slide(blank_slide_layout)
    bg11 = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg11.fill.solid()
    bg11.fill.fore_color.rgb = COLOR_DARK
    bg11.line.fill.background()

    t_box11 = slide11.shapes.add_textbox(Inches(1), Inches(1.2), Inches(11.333), Inches(5),)
    tf11 = t_box11.text_frame
    tf11.word_wrap = True

    pk1 = tf11.paragraphs[0]
    pk1.text = "TERIMA KASIH"
    pk1.font.name = "Arial"
    pk1.font.size = Pt(36)
    pk1.font.bold = True
    pk1.font.color.rgb = COLOR_ACCENT

    pk2 = tf11.add_paragraph()
    pk2.text = "Sistem Informasi Manajemen Beasiswa Terpadu (Beasiswa Blitar Mengabdi) V2.2.0 Siap Digunakan & Dideploy"
    pk2.font.name = "Arial"
    pk2.font.size = Pt(16)
    pk2.font.bold = True
    pk2.font.color.rgb = COLOR_WHITE
    pk2.space_before = Pt(10)

    pk3 = tf11.add_paragraph()
    pk3.text = "PUSAT LAYANAN INFORMASI & BANTUAN TEKNIS OFFICIAL:"
    pk3.font.name = "Arial"
    pk3.font.size = Pt(12)
    pk3.font.bold = True
    pk3.font.color.rgb = COLOR_ACCENT
    pk3.space_before = Pt(25)

    contacts = [
        "Instansi Pengelola: Dinas Kepemudaan dan Olahraga (Dispora) Kabupaten Blitar",
        "Penanggung Jawab Layanan: Bapak Akhyat",
        "Contact Person WhatsApp: 0813-3400-1600",
        "Alamat Kantor: Jl. Raya Sawahan Pojok, Kec. Garum, Kabupaten Blitar",
        "Jam Pelayanan Resmi: Senin — Jumat, 08:00 — 16:00 WIB"
    ]

    for c in contacts:
        pc = tf11.add_paragraph()
        pc.text = f"•  {c}"
        pc.font.name = "Arial"
        pc.font.size = Pt(11)
        pc.font.color.rgb = COLOR_WHITE
        pc.space_before = Pt(4)

    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    public_ppt = os.path.join(os.getcwd(), "public", "PRESENTASI_BEASISWA_BLITAR_MENGABDI.pptx")
    docs_ppt = os.path.join(os.getcwd(), "docs", "PRESENTASI_BEASISWA_BLITAR_MENGABDI.pptx")

    create_presentation(public_ppt)
    create_presentation(docs_ppt)
